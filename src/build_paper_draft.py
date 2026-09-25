"""Build the evidence-based review draft from existing results, without training."""
from pathlib import Path
import csv, json, sys, re, shutil, hashlib
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vendor'))
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/paper_draft_20260925'
FIG=OUT/'figures'; FIG.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(Path.home()/'.codex/skills/math-modeling/tools/docx/scripts'))
from equations import latex2omml
def equation_xml(s):
 for a,b in [(r'\arg\max',r'\argmax'),(r'\widetilde',r'\tilde'),(r'\widehat',r'\hat'),(r'\top','⊤'),(r'\lceil','⌈'),(r'\rceil','⌉'),(r'\setminus','∖'),(r'\varnothing','∅'),(r'\mid','|')]:s=s.replace(a,b)
 return etree.fromstring(latex2omml(s))
plt.rcParams.update({'font.sans-serif':['Microsoft YaHei'],'axes.unicode_minus':False,
 'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none',
 'axes.titleweight':'bold','figure.facecolor':'white','savefig.facecolor':'white'})
BLUE='#287C8E'; RED='#C65B4B'; GOLD='#BC923D'; DARK='#233B50'
sources={}
def read(rel):
 p=ROOT/rel;sources[rel]=hashlib.sha256(p.read_bytes()).hexdigest()
 return list(csv.DictReader(p.open(encoding='utf-8-sig')))
def save(name,fig):
 fig.savefig(FIG/(name+'.png'),dpi=220,bbox_inches='tight');fig.savefig(FIG/(name+'.svg'),bbox_inches='tight');plt.close(fig)
def boxes(name,labels,arrows,figsize=(10,4)):
 f,a=plt.subplots(figsize=figsize);a.set(xlim=(0,10),ylim=(0,5));a.axis('off')
 for x,y,w,h,txt,col in labels:
  a.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.05,rounding_size=.12',facecolor=col,edgecolor='#B6C7CC',lw=.8));a.text(x+w/2,y+h/2,txt,ha='center',va='center',fontsize=10,color=DARK)
 for start,end in arrows:a.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','color':DARK,'lw':1.3})
 if name=='model':
  a.add_patch(FancyBboxPatch((4.65,3.95),3.7,.85,boxstyle='round,pad=.05',facecolor='#F3ECD9',edgecolor='#829BA5',linestyle='--'))
  a.text(6.5,4.375,'单模态辅助头 ×3\n极性＋强度（仅训练）',ha='center',va='center',fontsize=10,color=DARK)
  a.annotate('',xy=(5.6,3.9),xytext=(5.6,3.35),arrowprops={'arrowstyle':'->','color':DARK,'linestyle':'--'})
 save(name,f)

boxes('route',[(.1,3.4,2.5,1,'附件1：原视频与转写','#EEF4F6'),(3.2,3.4,3,1,'自提特征 → 词时间聚合','#DAEBEC'),(7,3.4,2.8,1,'问题一：对齐与质量分析','#DAEBEC'),(.1,1.5,2.5,1,'附件2：官方训练/验证特征','#EEF4F6'),(3.2,1.5,3,1,'掩码分支编码 → 双任务预测','#E2EAF2'),(7,1.5,2.8,1,'问题二：附件3预测','#E2EAF2'),(3.2,.0,3,1,'固定参数 → 模态/局部扰动','#F3ECD9'),(7,.0,2.8,1,'问题三：附件4解释','#F3ECD9')],[((2.65,3.9),(3.15,3.9)),((6.25,3.9),(6.95,3.9)),((2.65,2),(3.15,2)),((6.25,2),(6.95,2)),((4.7,1.45),(4.7,1.05)),((6.25,.5),(6.95,.5))])
f,a=plt.subplots(figsize=(10,3));a.set(xlim=(0,6),ylim=(-.6,3.2),yticks=[0,1,2],yticklabels=['视频帧支撑区间','音频短时窗口','词时间区间'],xlabel='时间（示意）')
for lo,hi,label in [(.5,1.4,'词1'),(1.7,3.2,'词2'),(3.6,5.3,'词3')]:a.broken_barh([(lo,hi-lo)],(1.8,.4),facecolors=BLUE);a.text((lo+hi)/2,2.4,label,ha='center');a.axvspan(lo,hi,alpha=.05,color=BLUE)
for t in np.arange(.2,5.8,.18):a.broken_barh([(t,.23)],(.85,.25),facecolors=GOLD,alpha=.6)
for j,t in enumerate(np.arange(.2,5.8,.5)):a.broken_barh([(t,.45)],(-.15,.3),facecolors=RED if j in [3,7] else BLUE,alpha=.75)
a.text(5.8,.4,'红色：无效帧',ha='right',color=RED);save('alignment',f)
q=read('results/q1_quality_proxy_metrics.csv');f,axs=plt.subplots(1,3,figsize=(11,3.1))
for ax,key,title,col in [(axs[0],'word_alignment_coverage','词覆盖率',BLUE),(axs[1],'face_detection_rate','人脸帧检出率',GOLD)]:
 vals=[float(x[key]) for x in q];ax.hist(vals,bins=np.linspace(0,1,11),color=col,edgecolor='white');ax.set(title=title,xlabel='比例',ylabel='视频数',xlim=(0,1))
for status,col in [('PASS',BLUE),('REVIEW',RED)]:
 z=[x for x in q if x['status']==status];axs[2].scatter([float(x['word_alignment_coverage']) for x in z],[float(x['face_detection_rate']) for x in z],s=20,alpha=.6,label=f'{status} ({len(z)})',color=col)
axs[2].set(xlabel='词覆盖率',ylabel='人脸帧检出率');axs[2].legend(fontsize=8);f.tight_layout();save('q1_quality',f)
p=ROOT/'results/q1_review_figures/-egA8-b7-3M$_$26_review.png';shutil.copyfile(p,FIG/'q1_example.png');sources[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
f,axs=plt.subplots(1,3,figsize=(10,3));
for ax,vals,title in zip(axs,[[24,5],[71.61,91.00],[1562,2008]],['零人脸片段数 ↓','有效人脸帧比例（%）↑','有效视觉token数 ↑']):
 ax.bar(['基础版','裁剪候选'],vals,color=[BLUE,GOLD],width=.55);ax.set_title(title,fontsize=10);ax.set_ylim(0,max(vals)*1.25)
 for i,v in enumerate(vals):ax.text(i,v,str(v),ha='center',va='bottom')
f.tight_layout();save('q1_crop',f)
boxes('model',[(.1,3.6,1.5,.8,'文本 50×768','#DAEBEC'),(.1,2,1.5,.8,'音频 50×74','#F3ECD9'),(.1,.4,1.5,.8,'视觉 50×35','#F4E5E1'),(2.1,3.5,2,1,'投影96维\n1层Transformer','#DAEBEC'),(2.1,1.9,2,1,'投影96维\n1层Transformer','#F3ECD9'),(2.1,.3,2,1,'投影96维\n1层Transformer','#F4E5E1'),(4.65,1.6,2.05,1.7,'掩码池化\n3×96维表示\n＋3项覆盖率','#EEF4F6'),(7.25,1.8,1.15,1.3,'拼接\n共享层128','#EEF4F6'),(8.9,3.2,1,1,'极性\n3类','#DAEBEC'),(8.9,.6,1,1,'强度\n3 tanh','#F3ECD9')],[((1.65,4),(2.05,4)),((1.65,2.4),(2.05,2.4)),((1.65,.8),(2.05,.8)),((4.15,4),(4.6,2.9)),((4.15,2.4),(4.6,2.4)),((4.15,.8),(4.6,1.9)),((6.75,2.4),(7.2,2.4)),((8.45,2.4),(8.85,3.7)),((8.45,2.4),(8.85,1.1))],(11,4.1))
cand=read('results/q2_tonight_20260924/candidate_means.csv');labels=['无辅助基准','门控','学习率1e−4','辅助头（最终）','门控＋辅助头','多段遮蔽']
f,axs=plt.subplots(1,2,figsize=(10,3.6));idx=np.arange(6)
axs[0].errorbar([float(x['mean_selection_J']) for x in cand],idx,xerr=[float(x['sd_selection_J']) for x in cand],fmt='o',color=BLUE,capsize=3);axs[0].set(yticks=idx,yticklabels=labels,xlabel='综合分J ↓');axs[0].invert_yaxis();axs[0].grid(axis='x',alpha=.2)
for i,x in enumerate(cand):
 xx=float(x['mean_masked_mae']);yy=float(x['mean_masked_macro_f1']);axs[1].scatter(xx,yy,color=RED if i==3 else BLUE,s=45);axs[1].annotate(str(i+1),(xx,yy),xytext=(5,4),textcoords='offset points')
axs[1].set(xlabel='54场景平均MAE ↓',ylabel='54场景平均宏F1 ↑',title='编号与左图自上而下对应',ylim=(.576,.594));f.tight_layout();save('q2_compare',f)
conf=read('results/q2_final_c_uni_20260924/neutral_confusion.csv');mat=np.array([int(x['count']) for x in conf]).reshape(3,3);f,a=plt.subplots(figsize=(5,3.8));a.imshow(mat,cmap='Blues');a.set(xticks=range(3),yticks=range(3),xticklabels=['负面','中性','正面'],yticklabels=['负面','中性','正面'],xlabel='预测类别',ylabel='真实类别')
for i in range(3):
 for j in range(3):a.text(j,i,str(mat[i,j]),ha='center',va='center',color='white' if mat[i,j]>140 else DARK,fontsize=13)
save('confusion',f)
sc=[]
for seed in [20260923,20260924,20260925]:sc+=read(f'results/q2_c_aug_1layer_concat_eff_tonight_uni_20260924_seed_{seed}/scenario_metrics.csv')
f,axs=plt.subplots(1,3,figsize=(11,3.3))
for ax,key,groups,ticks in [(axs[0],'modalities',['T','A','V','TA','TV','AV'],['T','A','V','TA','TV','AV']),(axs[1],'ratio',['0.1','0.3','0.5'],['10%','30%','50%']),(axs[2],'position',['front','middle','back'],['前','中','后'])]:
 actual=set(x[key] for x in sc)
 if key=='position' and not set(groups)<=actual:groups=['front','middle','end'] if 'end' in actual else sorted(actual)
 vals=[np.mean([float(x['macro_f1']) for x in sc if x[key]==g]) for g in groups];ax.plot(ticks,vals,'o-',color=BLUE);ax.set_ylim(.54,.64);ax.set_ylabel('平均宏F1');ax.grid(alpha=.2)
axs[0].set_title('被遮模态组合');axs[1].set_title('名义遮蔽比例');axs[2].set_title('区间位置');f.tight_layout();save('missing',f)
w=read('results/q3_c_uni_validation_full2_20260925/valid_word_check.csv');top=np.array([float(x['original_word_recode_class_drop']) for x in w if x['kind']=='top']);rand=np.array([float(x['original_word_recode_class_drop']) for x in w if x['kind']=='random']);f,axs=plt.subplots(1,2,figsize=(9,3.5));axs[0].bar(['首选窗口','随机窗口'],[top.mean(),rand.mean()],color=[BLUE,GOLD],width=.5)
for i,v in enumerate([top.mean(),rand.mean()]):axs[0].text(i,v,f'{v:.4f}',ha='center',va='bottom')
axs[0].set(ylabel='平均类别概率下降',ylim=(0,.19));axs[1].scatter(rand,top,s=18,alpha=.65,color=BLUE);lo=min(top.min(),rand.min())-.03;hi=max(top.max(),rand.max())+.03;axs[1].plot([lo,hi],[lo,hi],'--',lw=1,color='gray');axs[1].set(xlabel='随机窗口概率下降',ylabel='首选窗口概率下降',xlim=(lo,hi),ylim=(lo,hi));f.tight_layout();save('faithfulness',f)
con=[x for x in read('results/q3_c_uni_final_20260925/modality_contributions.csv') if x['reference']=='mean'];f,a=plt.subplots(figsize=(10,3));mm=np.array([[float(x['class_'+m]) for x in con] for m in ['text','audio','vision']]);lim=np.abs(mm).max();im=a.imshow(mm,cmap='RdBu_r',vmin=-lim,vmax=lim,aspect='auto');a.set(xticks=range(20),xticklabels=[x['sample_id'] for x in con],yticks=range(3),yticklabels=['文本','音频','视觉'],xlabel='附件4样本');f.colorbar(im,ax=a,label='类别概率贡献');f.tight_layout();save('contributions',f)
ev=read('results/q3_c_uni_final_20260925/key_evidence.csv');f,axs=plt.subplots(1,3,figsize=(11,3));cc=next(x for x in con if x['sample_id']=='10');vals=[float(cc['class_'+m]) for m in ['text','audio','vision']];axs[0].barh(['文本','音频','视觉'],vals,color=[BLUE,RED,RED]);axs[0].axvline(0,color='gray',lw=.8);axs[0].set_title('样本10：类别贡献');axs[0].set_xlabel('Shapley贡献')
ee=[x for x in ev if x['sample_id']=='10' and x['modality']=='text'];axs[1].barh([x['text_span'] for x in ee],[float(x['original_text_recode_class_drop']) for x in ee],color=BLUE);axs[1].set_title('样本10：文本候选');axs[1].set_xlabel('原词重编码概率下降')
frame=ROOT/'results/q3_c_uni_final_20260925/frames/05_vision_1.jpg';axs[2].imshow(plt.imread(frame));axs[2].axis('off');axs[2].set_title('样本05：视觉候选关键帧\n约1.26—2.25秒（估计）',fontsize=10);f.tight_layout();save('cases',f)

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.25)
for st in ['Normal','Body Text']:
 s=doc.styles[st];s.font.name='Times New Roman';s.font.size=Pt(10.5);s.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');s.paragraph_format.line_spacing=1.25;s.paragraph_format.space_after=Pt(5)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12)]:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('233B50');s.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='基于可观测掩码与扰动归因的多模态情感预测'
head=sec.header.paragraphs[0];head.text='多模态情感预测 · 研究初稿';head.alignment=2;head.runs[0].font.size=Pt(8);head.runs[0].font.color.rgb=RGBColor(120,130,140)
foot=sec.footer.paragraphs[0];foot.alignment=1
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def para(txt,style=None):
 p=doc.add_paragraph(txt,style);p.paragraph_format.widow_control=True
 if style=='Caption' and txt.startswith('表'):p.paragraph_format.keep_with_next=True
 if not style:p.paragraph_format.first_line_indent=Pt(21)
 return p
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=True
 for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();x=OxmlElement('w:cantSplit');pr.append(x)
  if ri==0:
   rep=OxmlElement('w:tblHeader');pr.append(rep)
  for cell in row.cells:
   tcPr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','bottom']:
    z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single' if (ri==0 or (edge=='bottom' and ri==len(t.rows)-1)) else 'nil');z.set(qn('w:sz'),'8');z.set(qn('w:color'),'829BA5');b.append(z)
   tcPr.append(b)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF1F3');tcPr.append(sh)
   for p in cell.paragraphs:
    if len(rows)<=10:p.paragraph_format.keep_with_next=ri<len(t.rows)-1
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
    for run in p.runs:run.font.size=Pt(8 if len(headers)>6 else 9);run.bold=ri==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
def auto_table(tag):
 if tag=='@TABLE_CANDIDATES':
  para('表6 集中对照实验（三个种子的指标均值）','Caption');table(['配置','J ↓','原样宏F1','缺失宏F1','缺失MAE','改善种子'],[[labels[i]]+[f'{float(x[k]):.5f}' for k in ['mean_selection_J','mean_clean_macro_f1','mean_masked_macro_f1','mean_masked_mae']]+[x['improved_seed_count']+'/3' if i else '基准'] for i,x in enumerate(cand)])
 elif tag in ['@TABLE_A3','@TABLE_A4']:
  is3=tag.endswith('A3');rs=read('results/q2_attachment3_candidate_q2_c_aug_1layer_concat_eff_tonight_uni_20260924/predictions.csv' if is3 else 'results/q3_c_uni_final_20260925/predictions.csv');para('表8 附件3全量预测' if is3 else '表9 附件4全量预测','Caption')
  table(['样本','极性','强度','p负面','p中性','p正面'],[[x['sample_id'].replace('附件3_',''),{'negative':'负面','neutral':'中性','positive':'正面'}[x['polarity']]]+[f'{float(x[k]):.3f}' for k in ['strength','p_negative','p_neutral','p_positive']] for x in rs])
 elif tag=='@TABLE_Q1':
  inv=read('results/q1_feature_inventory.csv');qr={x['sample_id']:x for x in q};para('表A1 第一问100条样本基础版清单','Caption');table(['样本编号','状态','秒','词数','未解','音频有效','视觉有效'],[[x['sample_id'],x['status'],f"{float(qr[x['sample_id']]['video_duration_s']):.2f}",qr[x['sample_id']]['words'],x['unresolved_words'],x['audio_valid_tokens'],x['vision_valid_tokens']] for x in inv])
 elif tag=='@TABLE_EXPLANATIONS':
  preds=read('results/q3_c_uni_final_20260925/predictions.csv');rows=[]
  for x in preds:
   c=next(a for a in con if a['sample_id']==x['sample_id']);e=next((a for a in ev if a['sample_id']==x['sample_id'] and a['modality']==x['main_modality'] and a['rank']=='1'),None)
   desc=(e['text_span'] if e else '无正向候选');tm=(f"{float(e['start_s']):.2f}—{float(e['end_s']):.2f}s" if e and e['start_s'] else '未定位')
   if e and e['modality']=='text':tm+='；方向'+('一致' if float(e['original_text_recode_class_drop'])>0 else '不一致')
   rows.append([x['sample_id']]+[f'{float(c["class_"+m]):+.3f}' for m in ['text','audio','vision']]+[{'text':'文本','audio':'音频','vision':'视觉'}[x['main_modality']],desc+'\n'+tm])
  para('表B1 附件4类别贡献与主要模态首选候选','Caption');table(['样本','φ文本','φ音频','φ视觉','主模态','候选与定位'],rows)
lines=(OUT/'论文正文.md').read_text(encoding='utf-8').splitlines();i=0;eq=0;figs=0
while i<len(lines):
 l=lines[i].strip();i+=1
 if not l:continue
 if l.startswith('@TABLE_'):auto_table(l)
 elif l.startswith('$$ '):
  eq+=1;p=doc.add_paragraph();p.alignment=1;p.paragraph_format.space_after=Pt(8);p._p.append(equation_xml(l[3:-3]));p.add_run(f'    （{eq}）').font.size=Pt(9)
 elif l.startswith('!['):
  m=re.match(r'!\[(.*)\]\((.*)\)',l);p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(OUT/m[2]),width=Cm(16));cap=para(m[1],'Caption');cap.alignment=1;cap.paragraph_format.space_after=Pt(10);figs+=1
 elif l.startswith('|'):
  raw=[l]
  while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
  rr=[[c.strip() for c in s.strip('|').split('|')] for s in raw];table(rr[0],rr[2:])
 elif l.startswith('# '):
  title=l[2:]
  if title.startswith('基于'):p=para(title,'Title');p.alignment=1
  else:
   if title.startswith(('1 ','附录')):doc.add_page_break()
   para(title,'Heading 1')
 elif l.startswith('## '):para(l[3:],'Heading 2')
 else:para(l,'Caption' if re.match(r'^表\d',l) else None)
target=OUT/'多模态情感预测_研究初稿.docx';doc.save(target)
(OUT/'构建核验.json').write_text(json.dumps({'equations':eq,'figures':figs,'tables':len(doc.tables),'source_hashes':sources,'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'new_training':False,'q1_version':'q1 baseline; q1_v2 candidate only'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(target),'equations':eq,'figures':figs,'tables':len(doc.tables)},ensure_ascii=False))
