"""Build the evidence-based review draft from existing results, without training."""
from pathlib import Path
import csv, json, sys, re, shutil, hashlib
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vendor'))
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/paper_manuscript_v2_20260925'
FIG=OUT/'figures'; FIG.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(Path.home()/'.codex/skills/math-modeling/tools/docx/scripts'))
from equations import latex2omml
def equation_xml(s):
 for a,b in [(r'\arg\max',r'\argmax'),(r'\widetilde',r'\tilde'),(r'\widehat',r'\hat'),(r'\top','⊤'),(r'\lceil','⌈'),(r'\rceil','⌉'),(r'\setminus','∖'),(r'\varnothing','∅'),(r'\mid','|')]:s=s.replace(a,b)
 return etree.fromstring(latex2omml(s))
plt.rcParams.update({'font.sans-serif':['Microsoft YaHei'],'axes.unicode_minus':False,
 'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none',
 'axes.titleweight':'bold','figure.facecolor':'white','savefig.facecolor':'white'})
BLUE='#287C8E'; RED='#C65B4B'; GOLD='#BC923D'; DARK='#233B50'
sources={}
def read(rel):
 p=ROOT/rel;sources[rel]=hashlib.sha256(p.read_bytes()).hexdigest()
 return list(csv.DictReader(p.open(encoding='utf-8-sig')))
# Reuse the reviewed figures; no training or figure regeneration.
q=read('results/q1_quality_proxy_metrics.csv')
cand=read('results/q2_tonight_20260924/candidate_means.csv')
labels=['无辅助基准','门控','学习率1e−4','辅助头（最终）','门控＋辅助头','多段遮蔽']
con=[x for x in read('results/q3_c_uni_final_20260925/modality_contributions.csv') if x['reference']=='mean']
ev=read('results/q3_c_uni_final_20260925/key_evidence.csv')
for p in FIG.glob('*.png'):sources[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
sources[str((OUT/'论文正文.md').relative_to(ROOT))]=hashlib.sha256((OUT/'论文正文.md').read_bytes()).hexdigest()

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.25)
for st in ['Normal','Body Text']:
 s=doc.styles[st];s.font.name='Times New Roman';s.font.size=Pt(10.5);s.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');s.paragraph_format.line_spacing=1.25;s.paragraph_format.space_after=Pt(5)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12)]:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('233B50');s.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='基于可观测掩码与扰动归因的多模态情感预测'
from docx.enum.style import WD_STYLE_TYPE
for name in ['TOC 1','TOC 2']:
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Times New Roman';s.font.size=Pt(9);s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1.05
head=sec.header.paragraphs[0];head.text='多模态情感预测 · 论文修订稿';head.alignment=2;head.runs[0].font.size=Pt(8);head.runs[0].font.color.rgb=RGBColor(120,130,140)
foot=sec.footer.paragraphs[0];foot.alignment=1
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def para(txt,style=None):
 p=doc.add_paragraph(txt,style);p.paragraph_format.widow_control=True
 if style=='Caption' and txt.startswith('表'):p.paragraph_format.keep_with_next=True
 if not style:p.paragraph_format.first_line_indent=Pt(21)
 return p
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=True
 for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 if headers[0]=='样本编号':
  t.autofit=False
  for col,width in zip(t.columns,[4.1,1.6,1.3,1.3,1.3,3.45,3.45]):col.width=Cm(width)
  for row in t.rows:
   for cell,width in zip(row.cells,[4.1,1.6,1.3,1.3,1.3,3.45,3.45]):cell.width=Cm(width)
 if headers[-1]=='候选与定位':
  t.autofit=False
  for col,width in zip(t.columns,[1,1.5,1.5,1.5,1.5,8.5]):col.width=Cm(width)
  for row in t.rows:
   for cell,width in zip(row.cells,[1,1.5,1.5,1.5,1.5,8.5]):cell.width=Cm(width)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();x=OxmlElement('w:cantSplit');pr.append(x)
  if ri==0:
   rep=OxmlElement('w:tblHeader');pr.append(rep)
  for cell in row.cells:
   tcPr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','bottom']:
    z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single' if (ri==0 or (edge=='bottom' and ri==len(t.rows)-1)) else 'nil');z.set(qn('w:sz'),'8');z.set(qn('w:color'),'829BA5');b.append(z)
   tcPr.append(b)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF1F3');tcPr.append(sh)
   for p in cell.paragraphs:
    if len(rows)<=10:p.paragraph_format.keep_with_next=ri<len(t.rows)-1
    p.paragraph_format.space_after=Pt(2 if headers[-1]=='候选与定位' else 3);p.paragraph_format.space_before=Pt(2 if headers[-1]=='候选与定位' else 3);p.paragraph_format.line_spacing=1.05
    for run in p.runs:run.font.size=Pt(8 if len(headers)>6 else 9);run.bold=ri==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
def auto_table(tag):
 if tag=='@TABLE_CANDIDATES':
  para('表6 集中对照实验（三个种子的指标均值）','Caption');table(['配置','J ↓','原样宏F1','缺失宏F1','缺失MAE','改善种子'],[[labels[i]]+[f'{float(x[k]):.5f}' for k in ['mean_selection_J','mean_clean_macro_f1','mean_masked_macro_f1','mean_masked_mae']]+[x['improved_seed_count']+'/3' if i else '基准'] for i,x in enumerate(cand)])
 elif tag in ['@TABLE_A3','@TABLE_A4']:
  is3=tag.endswith('A3');rs=read('results/q2_attachment3_candidate_q2_c_aug_1layer_concat_eff_tonight_uni_20260924/predictions.csv' if is3 else 'results/q3_c_uni_final_20260925/predictions.csv');para('表8 附件3全量预测' if is3 else '表9 附件4全量预测','Caption')
  table(['样本','极性','强度','p负面','p中性','p正面'],[[x['sample_id'].replace('附件3_',''),{'negative':'负面','neutral':'中性','positive':'正面'}[x['polarity']]]+[f'{float(x[k]):.3f}' for k in ['strength','p_negative','p_neutral','p_positive']] for x in rs])
 elif tag=='@TABLE_Q1':
  inv=read('results/q1_feature_inventory.csv');qr={x['sample_id']:x for x in q};para('表A1 第一问100条样本基础版清单','Caption');table(['样本编号','状态','秒','词数','未解','音频有效','视觉有效'],[[x['sample_id'],x['status'],f"{float(qr[x['sample_id']]['video_duration_s']):.2f}",qr[x['sample_id']]['words'],x['unresolved_words'],x['audio_valid_tokens'],x['vision_valid_tokens']] for x in inv])
 elif tag=='@TABLE_EXPLANATIONS':
  preds=read('results/q3_c_uni_final_20260925/predictions.csv');rows=[]
  for x in preds:
   c=next(a for a in con if a['sample_id']==x['sample_id']);e=next((a for a in ev if a['sample_id']==x['sample_id'] and a['modality']==x['main_modality'] and a['rank']=='1'),None)
   desc=(e['text_span'] if e else '无正向候选');tm=(f"{float(e['start_s']):.2f}—{float(e['end_s']):.2f}s" if e and e['start_s'] else '未定位')
   if e and e['modality']=='text':tm+='；方向'+('一致' if float(e['original_text_recode_class_drop'])>0 else '不一致')
   rows.append([x['sample_id']]+[f'{float(c["class_"+m]):+.3f}' for m in ['text','audio','vision']]+[{'text':'文本','audio':'音频','vision':'视觉'}[x['main_modality']],desc+'\n'+tm])
  para('表B1 附件4类别贡献与主要模态首选候选','Caption');table(['样本','φ文本','φ音频','φ视觉','主模态','候选与定位'],rows)
lines=(OUT/'论文正文.md').read_text(encoding='utf-8').splitlines();i=0;eq=0;figs=0
while i<len(lines):
 l=lines[i].strip();i+=1
 if not l:continue
 if l.startswith('@TABLE_'):auto_table(l)
 elif l.startswith('$$ '):
  eq+=1;p=doc.add_paragraph();p.alignment=1;p.paragraph_format.space_after=Pt(8);p._p.append(equation_xml(l[3:-3]));p.add_run(f'    （{eq}）').font.size=Pt(9)
 elif l.startswith('!['):
  m=re.match(r'!\[(.*)\]\((.*)\)',l);p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(OUT/m[2]),width=Cm(9.5 if 'confusion' in m[2] else 14.8));cap=para(m[1],'Caption');cap.alignment=1;cap.paragraph_format.space_after=Pt(10);figs+=1
 elif l.startswith('|'):
  raw=[l]
  while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
  rr=[[c.strip() for c in s.strip('|').split('|')] for s in raw];table(rr[0],rr[2:])
 elif l.startswith('# '):
  title=l[2:]
  if title.startswith('基于'):p=para(title,'Title');p.alignment=1
  else:
   if title.startswith('1 '):
    doc.add_page_break()
    para('目录','Title')
    toc=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');toc._p.append(field)
    doc.add_page_break()
   elif title.startswith('附录') or title=='参考文献':doc.add_page_break()
   para(title,'Heading 1')
 elif l.startswith('## '):para(l[3:],'Heading 2')
 else:para(l,'Caption' if re.match(r'^表\d',l) else None)
target=OUT/'多模态情感预测_论文修订稿.docx';doc.save(target)
(OUT/'构建核验.json').write_text(json.dumps({'equations':eq,'figures':figs,'tables':len(doc.tables),'source_hashes':sources,'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'new_training':False,'q1_version':'q1 baseline; q1_v2 candidate only'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(target),'equations':eq,'figures':figs,'tables':len(doc.tables)},ensure_ascii=False))
